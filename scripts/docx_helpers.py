"""Small python-docx helper layer used by the report generator."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

PENDING = "Pending execution"
BLUE = RGBColor(0x1F, 0x4E, 0x8C)
GREY = RGBColor(0x52, 0x51, 0x4E)


def _shade(cell, hex_fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _page_number_footer(section) -> None:
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("RoadEye - ML Innovative Assignment   |   Page ")
    run.font.size = Pt(8)
    run.font.color.rgb = GREY
    r2 = p.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        r2._r.append(el)
    r2.font.size = Pt(8)


class Doc:
    """Thin wrapper: headings, paragraphs, bullets, tables, figures, callouts."""

    def __init__(self, title: str, subtitle: str = "", banner: str | None = None, meta_lines=None):
        self.d = Document()
        sec = self.d.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)  # A4
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin = sec.bottom_margin = Cm(2.0)
        st = self.d.styles["Normal"]
        st.font.name = "Calibri"
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        st.font.size = Pt(10.5)
        st.paragraph_format.space_after = Pt(6)
        for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
            h = self.d.styles[name]
            h.font.name = "Calibri"
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = BLUE
        _page_number_footer(sec)
        t = self.d.add_paragraph()
        t.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = t.add_run(title)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(22), BLUE
        if subtitle:
            s = self.d.add_paragraph()
            rs = s.add_run(subtitle)
            rs.font.size, rs.font.color.rgb = Pt(12), GREY
        for line in meta_lines or []:
            m = self.d.add_paragraph()
            rm = m.add_run(line)
            rm.font.size, rm.font.color.rgb = Pt(9.5), GREY
            m.paragraph_format.space_after = Pt(0)
        if banner:
            self.callout(banner, fill="FDE9D9")

    # ---- text ----
    def h(self, text: str, level: int = 1):
        return self.d.add_heading(text, level=min(level, 3))

    def p(self, text: str = "", bold: bool = False, italic: bool = False, size: float | None = None):
        para = self.d.add_paragraph()
        for i, chunk in enumerate(text.split("**")):  # **bold** inline markup
            if not chunk:
                continue
            r = para.add_run(chunk)
            r.bold, r.italic = bold or (i % 2 == 1), italic
            if size:
                r.font.size = Pt(size)
        return para

    def rich(self, parts: list[tuple[str, bool]]):
        para = self.d.add_paragraph()
        for text, bold in parts:
            para.add_run(text).bold = bold
        return para

    def bullets(self, items, style: str = "List Bullet"):
        for it in items:
            para = self.d.add_paragraph(style=style)
            if isinstance(it, tuple):  # (bold lead, rest)
                para.add_run(it[0]).bold = True
                para.add_run(it[1])
            else:
                para.add_run(str(it))
            para.paragraph_format.space_after = Pt(2)

    def numbered(self, items):
        self.bullets(items, style="List Number")

    def callout(self, text: str, fill: str = "EAF1FB"):
        tbl = self.d.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        _shade(cell, fill)
        cell.paragraphs[0].add_run(text).font.size = Pt(10)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    # ---- tables / figures ----
    def table(self, header: list[str], rows: list[list], col_widths_cm: list[float] | None = None,
              font_size: float = 9, caption: str | None = None):
        if caption:
            c = self.d.add_paragraph()
            rc = c.add_run(caption)
            rc.bold, rc.font.size = True, Pt(9.5)
            c.paragraph_format.space_after = Pt(2)
            c.paragraph_format.keep_with_next = True
        tbl = self.d.add_table(rows=1, cols=len(header))
        tbl.style = "Table Grid"
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        for i, name in enumerate(header):
            cell = tbl.rows[0].cells[i]
            _shade(cell, "1F4E8C")
            run = cell.paragraphs[0].add_run(str(name))
            run.bold, run.font.size, run.font.color.rgb = True, Pt(font_size), RGBColor(255, 255, 255)
        for ri, row in enumerate(rows):
            cells = tbl.add_row().cells
            for i, val in enumerate(row):
                run = cells[i].paragraphs[0].add_run("" if val is None else str(val))
                run.font.size = Pt(font_size)
                if ri % 2 == 1:
                    _shade(cells[i], "F3F6FA")
        if col_widths_cm:
            for row in tbl.rows:
                for i, w in enumerate(col_widths_cm):
                    row.cells[i].width = Cm(w)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)
        return tbl

    def df_table(self, df: pd.DataFrame, caption: str | None = None, font_size: float = 8.5, fmt=None):
        fmt = fmt or {}
        rows = [[(fmt[c](v) if c in fmt and pd.notna(v) else (f"{v:.4f}" if isinstance(v, float) else v))
                 for c, v in zip(df.columns, r)] for r in df.itertuples(index=False)]
        return self.table(list(df.columns), rows, caption=caption, font_size=font_size)

    def figure(self, path: Path | str | None, caption: str, width_cm: float = 15.5):
        path = Path(path) if path else None
        if path is not None and path.exists():
            self.d.add_picture(str(path), width=Cm(width_cm))
            self.d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            c = self.d.add_paragraph()
            c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = c.add_run(caption)
            r.italic, r.font.size, r.font.color.rgb = True, Pt(9), GREY
        else:
            self.callout(f"[Figure pending execution] {caption}  (generated by the pipeline as "
                         f"{path.name if path else 'a results figure'})", fill="F3F3F3")

    def page_break(self):
        self.d.add_page_break()

    def save(self, path: Path | str) -> Path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.d.save(path)
        return path
